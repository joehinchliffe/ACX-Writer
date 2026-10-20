"""Turn the Markdown that Claude writes into Word documents."""

from __future__ import annotations

import re
from pathlib import Path

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt
except ImportError:  # pragma: no cover - runtime dependency check
    raise SystemExit("Missing dependency: pip install -r requirements.txt")

_INLINE = re.compile(r"(\*\*[^*]+\*\*|<b>.*?</b>|\*[^*\s][^*]*\*|_[^_\s][^_]*_)", re.IGNORECASE)


def add_inline(paragraph, text: str):
    """Add text to a paragraph, turning **bold**, <b>bold</b> and *italic* into runs."""
    for part in _INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.lower().startswith("<b>"):
            paragraph.add_run(re.sub(r"</?b>", "", part, flags=re.IGNORECASE)).bold = True
        elif (part[0] in "*_") and part.endswith(part[0]) and len(part) > 2:
            paragraph.add_run(part[1:-1]).italic = True
        else:
            paragraph.add_run(part)


def add_markdown(doc, markdown: str, *, heading_offset: int = 0):
    """Append Markdown (headings, lists, tables, paragraphs) to a document."""
    lines = markdown.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()
        if not stripped or stripped in ("---", "***", "```"):
            i += 1
            continue
        if stripped.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                row = lines[i].strip()
                if not re.fullmatch(r"\|[\s\-:|]+\|?", row):
                    rows.append([c.strip() for c in row.strip("|").split("|")])
                i += 1
            _add_table(doc, rows)
            continue
        heading = re.match(r"^(#{1,6})\s+(.*)", stripped)
        if heading:
            level = min(len(heading.group(1)) + heading_offset, 9)
            doc.add_heading(_strip_md(heading.group(2)), level=level)
        elif re.match(r"^[-*•]\s+", stripped):
            add_inline(doc.add_paragraph(style="List Bullet"), re.sub(r"^[-*•]\s+", "", stripped))
        elif re.match(r"^\d+[.)]\s+", stripped):
            add_inline(doc.add_paragraph(style="List Number"), re.sub(r"^\d+[.)]\s+", "", stripped))
        elif stripped.startswith(">"):
            p = doc.add_paragraph(style="Quote")
            add_inline(p, stripped.lstrip("> "))
        else:
            add_inline(doc.add_paragraph(), stripped)
        i += 1


def _strip_md(text: str) -> str:
    return re.sub(r"[*_`]", "", text).strip()


def _add_table(doc, rows: list[list[str]]):
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=0, cols=cols)
    table.style = "Table Grid"
    for r, cells in enumerate(rows):
        row = table.add_row()
        for c in range(cols):
            p = row.cells[c].paragraphs[0]
            add_inline(p, cells[c] if c < len(cells) else "")
            if r == 0:
                for run in p.runs:
                    run.bold = True
                shd = OxmlElement("w:shd")
                shd.set(qn("w:fill"), "D9E1F2")
                shd.set(qn("w:val"), "clear")
                row.cells[c]._tc.get_or_add_tcPr().append(shd)
    doc.add_paragraph()


def save_markdown_docx(markdown: str, path: Path, title: str):
    """Save a research or marketing document: a title, then the Markdown body."""
    doc = Document()
    doc.add_heading(title, level=0)
    add_markdown(doc, markdown)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))


def save_manuscript_docx(book: dict, chapters: list[dict], path: Path):
    """Save the manuscript: title page, copyright page, then one chapter per page."""
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Georgia"
    style.font.size = Pt(12)

    for text, size, bold in ((book["title"], 28, True), (book["subtitle"], 16, False), (book["pen_name"], 14, False)):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.font.size = Pt(size)
        run.bold = bold
    doc.add_page_break()

    for para in book["copyright"].split("\n\n"):
        p = doc.add_paragraph()
        run = p.add_run(para.strip())
        run.font.size = Pt(10)
    doc.add_page_break()

    for n, chapter in enumerate(chapters):
        if n:
            doc.add_page_break()
        doc.add_heading(_strip_md(chapter["heading"]), level=1)
        add_markdown(doc, chapter["text"])

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))


def docx_text(path: Path) -> str:
    return "\n".join(p.text for p in Document(str(path)).paragraphs if p.text.strip())
