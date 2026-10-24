"""Steps 5-7: book descriptions, categories and keywords, cover brief and Ideogram prompts."""

from __future__ import annotations

import base64
import mimetypes
from pathlib import Path

from docx import Document

from . import llm
from .docs import add_inline, save_markdown_docx
from .project import PROMPTS_DIR, Book, render, word_count
from .research import reader_text

# Prepended to the audiobook versions of each description.
AUDIOBOOK_TAGLINE_HTML = "<b>🎧 Audiobook + PDF included — get both formats with one purchase.</b>"
AUDIOBOOK_TAGLINE_PLAIN = "🎧 AUDIOBOOK + PDF INCLUDED — get both formats with one purchase."
SPLIT = "---PLAIN TEXT---"


def split_versions(description: str) -> tuple[str, str]:
    html, _, plain = description.partition(SPLIT)
    return html.strip(), plain.strip()


def add_description(doc, label: str, html: str, plain: str):
    doc.add_heading(label, level=1)
    doc.add_heading("Amazon KDP version (with bold formatting)", level=2)
    for line in html.splitlines():
        if line.strip():
            add_inline(doc.add_paragraph(), line.strip())
    doc.add_heading("HTML to paste into KDP", level=2)
    doc.add_paragraph(html)
    if plain:
        doc.add_heading("Plain text version (Audible / ACX)", level=2)
        for line in plain.splitlines():
            if line.strip():
                doc.add_paragraph(line.strip())


def descriptions(book: Book, research: dict, manuscript: str):
    print("\n[5/7] Book descriptions")
    values = dict(title=book.title, subtitle=book.subtitle, audience=research["audience"],
                  pain_point=research["pain_point"], transformation=research["transformation"],
                  manuscript=manuscript)
    audience = book.step("description_audience.md", lambda: llm.ask(
        render("description_audience", reader_brief=reader_text(research), **values),
        kind="marketing", label="description v1"))
    hormozi = book.step("description_hormozi.md", lambda: llm.ask(
        render("description_hormozi", **values), kind="marketing", label="description v2"))

    doc = Document()
    doc.add_heading(f"Book Descriptions: {book.title}", level=0)
    for n, (name, text) in enumerate((("Version 1 — Audience-Optimised", audience),
                                      ("Version 2 — Alex Hormozi Style", hormozi))):
        html, plain = split_versions(text)
        if n:
            doc.add_page_break()
        add_description(doc, f"{name} (eBook)", html, plain)
        doc.add_page_break()
        add_description(doc, f"{name} (Audiobook)", f"{AUDIOBOOK_TAGLINE_HTML}\n\n{html}",
                        f"{AUDIOBOOK_TAGLINE_PLAIN}\n\n{plain}" if plain else "")
    doc.save(str(book.root / f"Descriptions - {book.file_title}.docx"))
    print("  ✓ Descriptions saved")


def categories(book: Book, research: dict, chapters: list[dict], *, web_search: bool = True):
    print("\n[6/7] Categories and keywords")
    text = book.step("categories.md", lambda: llm.ask(
        render("categories", title=book.title, subtitle=book.subtitle, audience=research["audience"],
               pain_point=research["pain_point"], transformation=research["transformation"],
               chapter_titles="\n".join(f"- {c['heading']}" for c in chapters),
               # ACX narration runs at roughly 155 words a minute.
               audio_hours=f"{sum(word_count(c['text']) for c in chapters) / 155 / 60:.1f}"),
        kind="marketing", web_search=web_search, label="categories"))
    save_markdown_docx(text, book.root / f"Categories - {book.file_title}.docx", f"Categories — {book.title}")
    print("  ✓ Categories saved")


def image_block(path: Path) -> dict:
    media_type = mimetypes.guess_type(path.name)[0] or "image/jpeg"
    return {"type": "image", "source": {"type": "base64", "media_type": media_type,
                                        "data": base64.standard_b64encode(path.read_bytes()).decode()}}


def cover(book: Book, research: dict, outline_text: str, competitor_covers: list[Path]):
    print("\n[7/7] Cover brief and Ideogram prompts")
    images = [image_block(p) for p in sorted((PROMPTS_DIR / "cover_types").glob("*.jpg"))]
    if not images:
        print("  ⚠ No Cover Types images in prompts/cover_types (page-1.jpg … page-6.jpg). "
              "The cover brief will be written without the examples.")
    note = ""
    if competitor_covers:
        images += [image_block(p) for p in competitor_covers]
        note = (f"The last {len(competitor_covers)} attached image(s) are competitor covers from this niche. "
                "Analyse them and make sure this cover stands apart from them.")

    brief = book.step("cover_brief.md", lambda: llm.ask(
        images + [{"type": "text", "text": render(
            "cover_brief", title=book.title, subtitle=book.subtitle, pen_name=book.pen_name,
            drp=research["drp"], outline=outline_text, competitor_note=note)}],
        kind="marketing", label="cover brief"))
    save_markdown_docx(brief, book.research / f"Cover Brief - {book.file_title}.docx", f"Cover Brief — {book.title}")

    prompts = book.step("ideogram_prompts.md", lambda: llm.ask(
        render("ideogram_prompts", cover_brief=brief, reader_brief=reader_text(research),
               title=book.title, subtitle=book.subtitle, pen_name=book.pen_name),
        kind="marketing", label="ideogram prompts"))
    save_markdown_docx(prompts, book.research / f"Ideogram Prompts - {book.file_title}.docx",
                       f"Ideogram Prompts — {book.title}")
    print("  ✓ Cover brief and Ideogram prompts saved (paste the prompts into ideogram.ai)")
